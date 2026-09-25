#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""青花瓷山藥糕：一頁步驟簡報（略過削皮、蒸熟）。"""

from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

ASSETS = Path(__file__).resolve().parent / "dessert-assets" / "ppt"
OUT = Path(__file__).resolve().parent / "青花瓷山藥糕製作步驟.pptx"

PINK = RGBColor(0xFF, 0x5C, 0x8A)
CORAL = RGBColor(0xFF, 0x6B, 0x6B)
ORANGE = RGBColor(0xFF, 0x9F, 0x43)
MINT = RGBColor(0x2E, 0xD4, 0xA0)
SKY = RGBColor(0x4D, 0xC1, 0xFF)
BLUE = RGBColor(0x5B, 0x8D, 0xFF)
LAVENDER = RGBColor(0xB8, 0x6B, 0xFF)
INK = RGBColor(0x3A, 0x2A, 0x42)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
QH_BG = RGBColor(0xEE, 0xF6, 0xFF)
FONT = "Hiragino Sans GB"

W = Inches(13.333)
H = Inches(7.5)

STEPS = [
    ("1", PINK, "qh_powder.jpg", "壓成細泥"),
    ("2", ORANGE, "qh_recipe.jpg", "加熟糕粉和糖"),
    ("3", LAVENDER, "qh_color.jpg", "一份染成藍色"),
    ("4", BLUE, "qh_mix.jpg", "輕輕捏出雲紋"),
    ("5", CORAL, "qh_fill.jpg", "包餡，搓成圓球"),
    ("6", MINT, "qh_whisk.jpg", "煮椰子水白涼粉"),
    ("7", SKY, "qh_pour.jpg", "先倒一層進模具"),
    ("8", BLUE, "qh_set.jpg", "放球，再倒滿"),
]


def _set_ea_font(run, name):
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rPr, qn(tag))
        el.set("typeface", name)


def set_run(run, text, size, color, bold=False):
    run.text = text
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.name = FONT
    run.font.italic = False
    _set_ea_font(run, FONT)


def add_text(slide, l, t, w, h, text, size=14, color=INK, bold=False,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    p.space_after = Pt(0)
    p.space_before = Pt(0)
    run = p.add_run()
    set_run(run, text, size, color, bold)
    return box


def solid(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def rect(slide, l, t, w, h, fill, rounded=False, adj=0.12):
    shp = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
    s = slide.shapes.add_shape(shp, l, t, w, h)
    solid(s, fill)
    if rounded:
        try:
            s.adjustments[0] = adj
        except Exception:
            pass
    s.line.fill.background()
    return s


def oval(slide, l, t, w, h, fill):
    s = slide.shapes.add_shape(MSO_SHAPE.OVAL, l, t, w, h)
    solid(s, fill)
    s.line.fill.background()
    return s


def photo(slide, name, l, t, w, h, border):
    pad = Inches(0.05)
    rect(slide, l, t, w, h, border, rounded=True, adj=0.08)
    rect(slide, l + pad, t + pad, w - pad * 2, h - pad * 2, WHITE, rounded=True, adj=0.07)
    slide.shapes.add_picture(
        str(ASSETS / name),
        l + pad + Inches(0.02), t + pad + Inches(0.02),
        w - pad * 2 - Inches(0.04), h - pad * 2 - Inches(0.04),
    )


def build():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    rect(slide, 0, 0, W, H, QH_BG)
    oval(slide, Inches(-0.7), Inches(-0.8), Inches(2.8), Inches(2.8), RGBColor(0xD6, 0xE8, 0xFF))
    oval(slide, Inches(11.8), Inches(-0.5), Inches(2.2), Inches(2.2), RGBColor(0xFF, 0xE3, 0xF0))
    oval(slide, Inches(12.3), Inches(6.1), Inches(1.7), Inches(1.7), RGBColor(0xFF, 0xEC, 0xB3))

    rect(slide, Inches(0.28), Inches(0.16), Inches(12.78), Inches(0.72), WHITE, rounded=True, adj=0.18)
    rect(slide, Inches(0.28), Inches(0.16), Inches(0.14), Inches(0.72), BLUE, rounded=True, adj=0.4)
    add_text(slide, Inches(0.58), Inches(0.16), Inches(8.2), Inches(0.28),
             "今天做這一味  ·  看圖就能做", 11, BLUE, True, PP_ALIGN.LEFT, MSO_ANCHOR.MIDDLE)
    add_text(slide, Inches(0.58), Inches(0.40), Inches(9.4), Inches(0.42),
             "青花瓷山藥糕｜製作步驟", 22, INK, True, PP_ALIGN.LEFT, MSO_ANCHOR.MIDDLE)
    add_text(slide, Inches(9.6), Inches(0.28), Inches(3.3), Inches(0.48),
             "冷藏後脫模就完成", 13, PINK, True, PP_ALIGN.RIGHT, MSO_ANCHOR.MIDDLE)

    cols, rows = 4, 2
    gap_x, gap_y = Inches(0.14), Inches(0.12)
    left, top = Inches(0.28), Inches(1.02)
    card_w = (Inches(12.78) - gap_x * (cols - 1)) / cols
    card_h = (Inches(6.22) - gap_y * (rows - 1)) / rows

    for i, (num, color, pic, title) in enumerate(STEPS):
        c, r = i % cols, i // cols
        x = left + c * (card_w + gap_x)
        y = top + r * (card_h + gap_y)
        rect(slide, x, y, card_w, card_h, WHITE, rounded=True, adj=0.08)
        photo(slide, pic, x + Inches(0.10), y + Inches(0.10),
              card_w - Inches(0.20), Inches(1.92), color)
        badge_w, badge_h = Inches(0.40), Inches(0.40)
        oval(slide, x + Inches(0.12), y + Inches(2.12), badge_w, badge_h, color)
        add_text(slide, x + Inches(0.12), y + Inches(2.12), badge_w, badge_h,
                 num, 14, WHITE, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
        add_text(slide, x + Inches(0.58), y + Inches(2.10),
                 card_w - Inches(0.70), Inches(0.46),
                 title, 16, INK, True, PP_ALIGN.LEFT, MSO_ANCHOR.MIDDLE)

    prs.save(OUT)
    print(f"saved {OUT}")


if __name__ == "__main__":
    build()
