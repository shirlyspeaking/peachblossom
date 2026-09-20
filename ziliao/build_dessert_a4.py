#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""兩份可編輯的一頁 A4 講義：材料 + 步驟，只放一張成品圖。"""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn, nsmap
from docx.shared import Cm, Emu, Pt, RGBColor, Twips

ROOT = Path(__file__).resolve().parent.parent
IMG = ROOT / "dessert-recipes" / "images"
OUT = Path(__file__).resolve().parent
FONT = "Hiragino Sans GB"


def set_run(run, text, size=11, bold=False, color=None, font=FONT):
    run.text = text
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = font
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), font)
    rFonts.set(qn("w:hAnsi"), font)
    rFonts.set(qn("w:eastAsia"), font)
    if color:
        run.font.color.rgb = RGBColor(*color)


def p_in(cell, text, size=11, bold=False, color=None, align="left", after=2, before=0):
    p = cell.paragraphs[0] if not cell.paragraphs[0].text and len(cell.paragraphs) == 1 else cell.add_paragraph()
    if cell.paragraphs[0] is p and cell.paragraphs[0].text:
        p = cell.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.08
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    if align == "center":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == "right":
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run()
    set_run(run, text, size, bold, color)
    return p


def first_p(cell, text, size=11, bold=False, color=None, align="left", after=2):
    p = cell.paragraphs[0]
    p.clear()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.08
    if align == "center":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    set_run(run, text, size, bold, color)
    return p


def shade(cell, hex_color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    for old in tcPr.findall(qn("w:shd")):
        tcPr.remove(old)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def borders(cell, color="FFFFFF", sz="4"):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    for old in tcPr.findall(qn("w:tcBorders")):
        tcPr.remove(old)
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), sz)
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        tcBorders.append(el)
    tcPr.append(tcBorders)


def no_borders(table):
    tbl = table._tbl
    tblPr = tbl.tblPr if tbl.tblPr is not None else OxmlElement("w:tblPr")
    for old in tblPr.findall(qn("w:tblBorders")):
        tblPr.remove(old)
    borders_el = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "nil")
        borders_el.append(el)
    tblPr.append(borders_el)


def set_cell_margin(cell, top=40, bottom=40, left=80, right=80):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    for old in tcPr.findall(qn("w:tcMar")):
        tcPr.remove(old)
    tcMar = OxmlElement("w:tcMar")
    for name, val in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        node = OxmlElement(f"w:{name}")
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")
        tcMar.append(node)
    tcPr.append(tcMar)


def to_emu(w):
    return w if hasattr(w, "twips") else Emu(int(w))


def set_table_width(table, width):
    w = to_emu(width)
    tbl = table._tbl
    tblPr = tbl.tblPr
    tblW = tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = OxmlElement("w:tblW")
        tblPr.append(tblW)
    tblW.set(qn("w:w"), str(int(w.twips)))
    tblW.set(qn("w:type"), "dxa")
    layout = tblPr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tblPr.append(layout)
    layout.set(qn("w:type"), "fixed")


def set_col_widths(table, widths):
    table.autofit = False
    table.allow_autofit = False
    widths = [to_emu(w) for w in widths]
    tbl = table._tbl
    tblGrid = tbl.find(qn("w:tblGrid"))
    if tblGrid is not None:
        tbl.remove(tblGrid)
    tblGrid = OxmlElement("w:tblGrid")
    for w in widths:
        gridCol = OxmlElement("w:gridCol")
        gridCol.set(qn("w:w"), str(int(w.twips)))
        tblGrid.append(gridCol)
    tbl.insert(0, tblGrid)
    for row in table.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = w


def prevent_break(table):
    for row in table.rows:
        tr = row._tr
        trPr = tr.get_or_add_trPr()
        cant = OxmlElement("w:cantSplit")
        trPr.append(cant)


def add_picture_cell(cell, path, width):
    p = cell.paragraphs[0]
    p.clear()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run()
    run.add_picture(str(path), width=width)


def build(title, kicker, accent, accent_hex, soft_hex, hero, materials, extra, steps, note):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(1.15)
    sec.bottom_margin = Cm(1.15)
    sec.left_margin = Cm(1.25)
    sec.right_margin = Cm(1.25)
    usable = to_emu(sec.page_width - sec.left_margin - sec.right_margin)
    left_w = Cm(11.2)
    right_w = to_emu(usable - left_w)

    n_rows = 5 + len(steps) + 1
    table = doc.add_table(rows=n_rows, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_col_widths(table, [left_w, right_w])
    set_table_width(table, usable)
    no_borders(table)

    def merge_row(r):
        table.cell(r, 0).merge(table.cell(r, 1))
        return table.cell(r, 0)

    title_cell = merge_row(0)
    shade(title_cell, accent_hex)
    set_cell_margin(title_cell, 90, 90, 140, 140)
    first_p(title_cell, kicker, 10, True, (255, 255, 255), after=0)
    p_in(title_cell, title, 22, True, (255, 255, 255), after=2)
    p_in(title_cell, "家政課一頁講義  ·  材料＋步驟  ·  看完就能做", 10, False, (255, 255, 255), after=2)

    bar = merge_row(1)
    shade(bar, accent_hex)
    set_cell_margin(bar, 50, 50, 120, 120)
    first_p(bar, "一、所需要的材料", 13, True, (255, 255, 255), after=0)

    mat = table.cell(2, 0)
    pic = table.cell(2, 1)
    shade(mat, soft_hex)
    shade(pic, soft_hex)
    set_cell_margin(mat, 80, 80, 120, 80)
    set_cell_margin(pic, 80, 50, 50, 80)
    mat.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    pic.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

    first_p(mat, "", 1, after=0)
    mat.paragraphs[0].clear()
    for i, (name, amt) in enumerate(materials):
        p = mat.paragraphs[0] if i == 0 else mat.add_paragraph()
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.05
        r1 = p.add_run()
        set_run(r1, f"•  {name}", 12, True, accent)
        r2 = p.add_run()
        set_run(r2, f"　　{amt}", 12, True, (58, 42, 66))

    add_picture_cell(pic, IMG / hero, Cm(7.0))
    cap = pic.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_before = Pt(2)
    cap.paragraph_format.space_after = Pt(0)
    cr = cap.add_run()
    set_run(cr, "成品參考", 9, True, accent)

    tip = merge_row(3)
    shade(tip, soft_hex)
    set_cell_margin(tip, 40, 80, 120, 120)
    first_p(tip, extra or "", 10, False, (90, 80, 96), after=0)

    h = merge_row(4)
    shade(h, accent_hex)
    set_cell_margin(h, 50, 50, 120, 120)
    first_p(h, "二、製作步驟", 13, True, (255, 255, 255), after=0)

    colors = [
        (255, 92, 138),
        (255, 159, 67),
        (46, 212, 160),
        (77, 193, 255),
        (91, 141, 255),
        (184, 107, 255),
        (255, 107, 107),
        (255, 208, 59),
    ]
    row_bg = ["FFFFFF", soft_hex]
    for i, text in enumerate(steps, start=1):
        cell = merge_row(4 + i)
        shade(cell, row_bg[(i - 1) % 2])
        set_cell_margin(cell, 50, 50, 120, 120)
        p = first_p(cell, "", 1, after=0)
        p.clear()
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.line_spacing = 1.08
        n = p.add_run()
        set_run(n, f"{i}.  ", 12, True, colors[(i - 1) % len(colors)])
        t = p.add_run()
        set_run(t, text, 12, False, (58, 42, 66))

    foot = merge_row(5 + len(steps))
    shade(foot, "FFFFFF")
    set_cell_margin(foot, 80, 40, 120, 120)
    first_p(foot, note, 9, False, (122, 100, 120), after=0)

    prevent_break(table)
    return doc


def main():
    qh = build(
        title="青花瓷山藥糕",
        kicker="中式甜點動手做",
        accent=(91, 141, 255),
        accent_hex="5B8DFF",
        soft_hex="EEF6FF",
        hero="qh_done.jpg",
        materials=[
            ("鐵棍山藥", "320 g"),
            ("熟糯米粉", "25 g"),
            ("白砂糖", "25 g"),
            ("茯苓粉", "10 g"),
            ("百合粉", "10 g"),
            ("蝶豆花粉", "適量"),
            ("椰子水 + 白涼粉", "約 10：1"),
        ],
        extra="小提醒：椰子水 150 毫升配白涼粉 15 克。可包喜歡的餡，也可以不包。山藥請戴手套。",
        steps=[
            "選鐵棍山藥，戴手套削皮。",
            "切成小塊，放入蒸烤箱，蒸到筷子能輕易插入。",
            "趁熱壓成細泥，壓到沒有顆粒。",
            "加入熟糯米粉 25g、白砂糖 25g、茯苓粉 10g、百合粉 10g，揉成光滑麵團。",
            "麵團分成兩份。其中一份加少量蝶豆花粉，揉成藍色。",
            "藍白兩色輕輕捏在一起（不要揉太勻，才有雲紋）。可包餡，再搓成圓球。",
            "椰子水加白涼粉攪勻，煮沸。",
            "先倒一層進模具，稍稍定型後放入山藥團，再倒滿。冷藏定型後脫模。",
        ],
        note="課堂小叮嚀：刀子請小心、熱的東西會燙、白涼粉一定要煮沸才會凝固。",
    )
    qh_path = OUT / "青花瓷山藥糕一頁講義.docx"
    qh.save(qh_path)

    xl = build(
        title="雪梨菊花銀耳凍",
        kicker="中式甜點動手做",
        accent=(255, 159, 67),
        accent_hex="FF9F43",
        soft_hex="FFF6E8",
        hero="xl_stack.jpg",
        materials=[
            ("雪梨", "去皮切丁"),
            ("銀耳", "適量"),
            ("冰糖", "適量"),
            ("菊花", "泡茶＋入模各用"),
            ("雪梨銀耳湯 + 白涼粉", "200 g + 20 g"),
            ("菊花茶 + 白涼粉", "100 g + 10 g"),
            ("泡好的菊花", "1 朵"),
        ],
        extra="好記口訣：湯／茶 : 白涼粉 ＝ 10 : 1。這樣凍起來 Q 彈，又不會太硬。",
        steps=[
            "先用熱水泡一壺菊花茶，備用。",
            "雪梨削皮、去核，切成均勻小丁。",
            "雪梨丁放進熱水，加入銀耳，蓋上蓋子悶煮大約半小時。",
            "銀耳軟了以後，加入冰糖攪到溶化並煮沸。",
            "舀出雪梨銀耳湯 200 克，加入白涼粉 20 克，攪勻煮沸。",
            "趁熱舀進花形模具（連同雪梨塊和銀耳），送進冰箱冷藏定型。",
            "另取菊花茶 100 克，加白涼粉 10 克，攪勻煮沸。",
            "倒進圓球模具，中間放一朵泡好的菊花，再補滿。冷藏後疊在雪梨凍上。",
        ],
        note="課堂小叮嚀：刀子請小心、熱湯會燙、白涼粉一定要煮沸才會凝固。",
    )
    xl_path = OUT / "雪梨菊花銀耳凍一頁講義.docx"
    xl.save(xl_path)
    print(f"saved {qh_path}")
    print(f"saved {xl_path}")


if __name__ == "__main__":
    main()
