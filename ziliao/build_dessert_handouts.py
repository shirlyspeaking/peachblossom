#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""產出兩份可發給學生的 Word 圖文步驟講義。"""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
IMG = ROOT / "dessert-recipes" / "images"
OUT = Path(__file__).resolve().parent


def set_run_font(run, name="微軟正黑體", size=12, bold=False, color=None):
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = name
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), name)
    rFonts.set(qn("w:hAnsi"), name)
    rFonts.set(qn("w:eastAsia"), name)
    if color:
        run.font.color.rgb = RGBColor(*color)


def add_p(doc, text, size=12, bold=False, color=None, align="left", space_after=8):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    if align == "center":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color)
    return p


def add_pic(doc, name, width=Cm(14.5)):
    path = IMG / name
    if not path.exists():
        raise FileNotFoundError(path)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(10)
    run = p.add_run()
    run.add_picture(str(path), width=width)


def shade_p(paragraph, hex_color):
    pPr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hex_color)
    pPr.append(shd)


def build_doc(title, kicker, lead, hero, materials_title, materials, extra, steps, finish_title, finish_imgs, finish_tip, accent):
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(1.6)
    section.bottom_margin = Cm(1.6)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)

    k = add_p(doc, kicker, 12, True, accent, space_after=4)
    shade_p(k, "FFF4F8")
    add_p(doc, title, 26, True, accent, space_after=6)
    add_p(doc, lead, 13, False, (90, 80, 96), space_after=10)
    add_pic(doc, hero)

    add_p(doc, "動手前先記住", 16, True, accent, space_after=6)
    add_p(doc, "刀子請小心；熱的東西會燙；白涼粉一定要煮沸才會凝固。山藥請戴手套。", 12, space_after=12)

    add_p(doc, materials_title, 16, True, accent, space_after=6)
    for name, amt, note in materials:
        add_p(doc, f"• {name}　{amt}　（{note}）", 13, space_after=2)
    if extra:
        t = add_p(doc, extra, 12, True, space_after=12)
        shade_p(t, "FFF3D1")

    add_p(doc, "製作步驟", 16, True, accent, space_after=8)
    for num, step_title, photo, body in steps:
        add_p(doc, f"步驟 {num}　{step_title}", 15, True, accent, space_after=4)
        add_pic(doc, photo)
        add_p(doc, body, 13, space_after=14)

    add_p(doc, finish_title, 16, True, accent, space_after=6)
    for name in finish_imgs:
        add_pic(doc, name)
    add_p(doc, finish_tip, 13, space_after=10)
    add_p(doc, "畫面來自原影片截圖。白涼粉記得煮沸、模具記得冷藏。做好記得拍照！", 11, False, (122, 100, 120), "center")
    return doc


def main():
    qh = build_doc(
        title="青花瓷山藥糕",
        kicker="給同學的圖文講義 · 第一款",
        lead="選鐵棍山藥、蒸熟壓泥、染成藍白雲紋，再包一層透明凍。看圖就能做。",
        hero="qh_done.jpg",
        materials_title="先把材料排好",
        materials=[
            ("鐵棍山藥", "320 g", "口感更粉糯"),
            ("熟糯米粉", "25 g", "讓它能搓成團"),
            ("白砂糖", "25 g", "甜度剛剛好"),
            ("茯苓粉", "10 g", "一起揉進山藥泥"),
            ("百合粉", "10 g", "一起揉進山藥泥"),
            ("蝶豆花粉", "適量", "染出青花藍"),
        ],
        extra="透明外衣：椰子水 + 白涼粉，大約 10 : 1。例如 150 毫升椰子水配 15 克白涼粉。可包喜歡的餡，也可以不包，直接搓球。",
        steps=[
            ("1", "選鐵棍山藥，削皮", "qh_peel.jpg", "做糕點要選鐵棍山藥，口感更好。戴手套削皮，黏液比較不會弄到手。"),
            ("2", "放進蒸烤箱蒸熟", "qh_steam.jpg", "去皮後切段或切小塊，放入蒸烤箱，蒸到筷子能輕易插入。熱熱的更好壓泥。"),
            ("3", "趁熱壓成細泥", "qh_mash.jpg", "蒸熟取出，用壓泥器壓到細細滑滑、沒有顆粒。越細，成品越像瓷器。"),
            ("4", "加粉揉成團", "qh_knead.jpg", "加入熟糯米粉 25g、白砂糖 25g、茯苓粉 10g、百合粉 10g，揉到不黏手的光滑麵團。"),
            ("5", "一份加蝶豆花粉", "qh_blue.jpg", "麵團分成兩份。其中一份加少量蝶豆花粉，揉成漂亮的藍色。粉少少加，顏色才清透。"),
            ("6", "雙色混合，可包餡", "qh_fill.jpg", "藍白兩色輕輕捏在一起，不要揉太勻，才有雲紋。可包喜歡的餡，再搓成圓球。"),
            ("7", "椰子水加白涼粉", "qh_whisk.jpg", "椰子水加白涼粉攪勻，煮沸。比例大約 10：1，例如 150 毫升椰子水配 15 克白涼粉。"),
            ("8", "倒模 → 放山藥團 → 再倒滿", "qh_set.jpg", "先倒一層進模具，稍稍定型後放入山藥團，再倒滿。冷藏定型，脫模就完成了！"),
        ],
        finish_title="完成！可以吃的小青花瓷",
        finish_imgs=["qh_done.jpg"],
        finish_tip="成功密技：藍白不要揉太開，紋路才像青花。外衣要等底部稍凝，球才不會沉到底。脫模前再冷藏一下，花紋更清楚。",
        accent=(91, 141, 255),
    )
    qh_path = OUT / "青花瓷山藥糕製作步驟.docx"
    qh.save(qh_path)

    xl = build_doc(
        title="雪梨菊花銀耳凍",
        kicker="給同學的圖文講義 · 第二款",
        lead="下面是雪梨銀耳凍，上面再放一顆藏著菊花的水晶球。清甜、潤喉、超漂亮。",
        hero="xl_stack.jpg",
        materials_title="分成兩層來做",
        materials=[
            ("雪梨銀耳湯", "200 g", "先把雪梨、銀耳、冰糖煮成湯"),
            ("白涼粉", "20 g", "做成雪梨銀耳凍"),
            ("菊花茶", "100 g", "另做上層水晶球"),
            ("白涼粉", "10 g", "做成菊花凍"),
            ("泡好的菊花", "1 朵", "藏進圓球中間"),
        ],
        extra="好記口訣：湯／茶 : 白涼粉 ＝ 10 : 1。這樣凍起來 Q 彈，又不會太硬。",
        steps=[
            ("1", "先泡一壺菊花茶", "xl_bloom.jpg", "準備菊花，熱水沖泡備用。待會一部分拿來做凍，一朵漂亮的花要藏進圓球裡。"),
            ("2", "雪梨去皮切小塊", "xl_cut.jpg", "雪梨削皮、去核，切成均勻小丁。塊不要太大，煮起來更快，凍裡也比較好看。"),
            ("3", "下鍋，加銀耳悶煮", "xl_fungus.jpg", "雪梨丁放進熱水裡，再加入銀耳。蓋上蓋子悶煮大約半小時，讓銀耳出膠、湯變黏潤。"),
            ("4", "最後加冰糖煮沸", "xl_sugar.jpg", "銀耳軟了以後，加入冰糖攪到溶化並煮沸。這鍋湯又香又甜，先聞一下超幸福。"),
            ("5", "取湯加白涼粉", "xl_powder.jpg", "舀出雪梨銀耳湯 200 克，加入白涼粉 20 克，攪勻煮沸。粉要完全化開，才不會一顆顆。"),
            ("6", "倒入模具冷藏", "xl_mold.jpg", "煮沸後趁熱舀進花形模具，連同雪梨塊和銀耳一起放。送進冰箱冷藏定型。"),
            ("7", "菊花茶加白涼粉煮沸", "xl_jupowder.jpg", "另起一鍋，倒入泡好的菊花茶 100 克，加白涼粉 10 克，攪勻煮沸。顏色會淡淡金黃。"),
            ("8", "倒模，中間放一朵花", "xl_flowerin.jpg", "倒進圓球模具，用夾子放進一朵泡好的菊花，再補滿茶凍。冷藏定型就完成了！"),
        ],
        finish_title="疊在一起，像一朵會融化的秋日",
        finish_imgs=["xl_base.jpg", "xl_stack.jpg"],
        finish_tip="怎麼擺盤？脫模後，先放上雪梨銀耳凍當底座，再把菊花水晶球輕輕放上去。想更漂亮，可以撒一點菊花碎。",
        accent=(255, 159, 67),
    )
    xl_path = OUT / "雪梨菊花銀耳凍製作流程.docx"
    xl.save(xl_path)
    print(f"saved {qh_path}")
    print(f"saved {xl_path}")


if __name__ == "__main__":
    main()
