#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""中式甜點教學簡報：青花瓷山藥糕 + 雪梨菊花銀耳凍。"""

from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

ASSETS = Path(__file__).resolve().parent / "dessert-assets" / "ppt"
OUT = Path(__file__).resolve().parent / "中式甜點教學簡報.pptx"

# 多巴胺：粉、橘、檸檬、薄荷、天藍、葡萄紫
PINK = RGBColor(0xFF, 0x5C, 0x8A)
CORAL = RGBColor(0xFF, 0x6B, 0x6B)
ORANGE = RGBColor(0xFF, 0x9F, 0x43)
YELLOW = RGBColor(0xFF, 0xD0, 0x3B)
MINT = RGBColor(0x2E, 0xD4, 0xA0)
SKY = RGBColor(0x4D, 0xC1, 0xFF)
BLUE = RGBColor(0x5B, 0x8D, 0xFF)
LAVENDER = RGBColor(0xB8, 0x6B, 0xFF)
INK = RGBColor(0x3A, 0x2A, 0x42)
INK_SOFT = RGBColor(0x7A, 0x64, 0x78)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
CREAM = RGBColor(0xFF, 0xFB, 0xF5)
BG = RGBColor(0xFF, 0xF4, 0xF8)
CARD = RGBColor(0xFF, 0xFF, 0xFF)
QH_BG = RGBColor(0xEE, 0xF6, 0xFF)
XL_BG = RGBColor(0xFF, 0xF8, 0xE8)

FONT_TITLE = "Hiragino Sans GB"
FONT_BODY = "Hiragino Sans GB"

W = Inches(13.333)
H = Inches(7.5)


def _set_ea_font(run, name):
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rPr, qn(tag))
        el.set("typeface", name)


def set_run(run, text, size, color, bold=False, font=FONT_BODY):
    run.text = text
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.name = font
    run.font.italic = False
    _set_ea_font(run, font)


def add_textbox(slide, l, t, w, h, text, size=14, color=INK, bold=False,
                font=FONT_BODY, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    p.space_after = Pt(0)
    p.space_before = Pt(0)
    p.line_spacing = 1.15
    run = p.add_run()
    set_run(run, text, size, color, bold, font)
    return box


def write_para(tf, text, size=14, color=INK, bold=False, font=FONT_BODY,
               align=PP_ALIGN.LEFT, space_after=6, first=False):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.space_after = Pt(space_after)
    p.space_before = Pt(0)
    p.line_spacing = 1.18
    run = p.add_run()
    set_run(run, text, size, color, bold, font)
    return p


def solid(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def stroke(shape, color, pt=1.0):
    shape.line.color.rgb = color
    shape.line.width = Pt(pt)


def rect(slide, l, t, w, h, fill, line=None, line_pt=1.25, rounded=False, adj=0.12):
    shp = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
    s = slide.shapes.add_shape(shp, l, t, w, h)
    solid(s, fill)
    if rounded:
        try:
            s.adjustments[0] = adj
        except Exception:
            pass
    if line:
        stroke(s, line, line_pt)
    else:
        s.line.fill.background()
    return s


def oval(slide, l, t, w, h, fill, line=None):
    s = slide.shapes.add_shape(MSO_SHAPE.OVAL, l, t, w, h)
    solid(s, fill)
    if line:
        stroke(s, line, 1.25)
    else:
        s.line.fill.background()
    return s


def star(slide, l, t, w, h, fill):
    s = slide.shapes.add_shape(MSO_SHAPE.STAR_5_POINT, l, t, w, h)
    solid(s, fill)
    s.line.fill.background()
    return s


def img(name):
    p = ASSETS / name
    if not p.exists():
        raise FileNotFoundError(p)
    return str(p)


def photo(slide, name, l, t, w, h, border=PINK):
    """彩色圓角外框 + 16:9 截圖。"""
    pad = Inches(0.07)
    frame = rect(slide, l, t, w, h, border, rounded=True, adj=0.08)
    rect(slide, l + pad, t + pad, w - pad * 2, h - pad * 2, WHITE, rounded=True, adj=0.07)
    slide.shapes.add_picture(
        img(name), l + pad + Inches(0.03), t + pad + Inches(0.03),
        w - pad * 2 - Inches(0.06), h - pad * 2 - Inches(0.06),
    )
    return frame


def badge(slide, l, t, text, fill, w=Inches(1.15), h=Inches(0.42)):
    rect(slide, l, t, w, h, fill, rounded=True, adj=0.5)
    add_textbox(slide, l, t, w, h, text, 13, WHITE, True, FONT_TITLE,
                PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)


def pill(slide, l, t, w, h, text, fill, fg=INK):
    rect(slide, l, t, w, h, fill, rounded=True, adj=0.5)
    add_textbox(slide, l, t, w, h, text, 13, fg, True, FONT_BODY,
                PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)


def add_bg(slide, color=BG):
    rect(slide, 0, 0, W, H, color)
    oval(slide, Inches(-0.8), Inches(-0.9), Inches(3.2), Inches(3.2), RGBColor(0xFF, 0xD6, 0xE8))
    oval(slide, Inches(11.6), Inches(-0.6), Inches(2.6), Inches(2.6), RGBColor(0xD6, 0xF4, 0xFF))
    oval(slide, Inches(12.2), Inches(5.8), Inches(2.0), Inches(2.0), RGBColor(0xFF, 0xEC, 0xB3))
    oval(slide, Inches(-0.5), Inches(6.4), Inches(1.8), Inches(1.8), RGBColor(0xD8, 0xF8, 0xE8))
    star(slide, Inches(12.55), Inches(0.95), Inches(0.38), Inches(0.38), YELLOW)
    star(slide, Inches(0.28), Inches(6.55), Inches(0.32), Inches(0.32), PINK)


def add_header(slide, title, accent=PINK, kicker="今天的甜點課"):
    rect(slide, Inches(0.42), Inches(0.22), Inches(12.5), Inches(0.78), WHITE, rounded=True, adj=0.18)
    rect(slide, Inches(0.42), Inches(0.22), Inches(0.16), Inches(0.78), accent, rounded=True, adj=0.4)
    add_textbox(slide, Inches(0.78), Inches(0.22), Inches(8.8), Inches(0.32),
                kicker, 11, accent, True, FONT_BODY, PP_ALIGN.LEFT, MSO_ANCHOR.MIDDLE)
    add_textbox(slide, Inches(0.78), Inches(0.46), Inches(11.8), Inches(0.48),
                title, 24, INK, True, FONT_TITLE, PP_ALIGN.LEFT, MSO_ANCHOR.MIDDLE)


def add_footer(slide, n, total, note="中式甜點動手做  ·  給國中同學的快樂食譜"):
    add_textbox(slide, Inches(0.50), Inches(7.14), Inches(9.6), Inches(0.26),
                note, 10, INK_SOFT, False, FONT_BODY)
    pill(slide, Inches(11.35), Inches(7.12), Inches(1.45), Inches(0.28),
         f"{n} / {total}", PINK, WHITE)


def new_slide(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def step_pair(slide, left, right):
    """兩個步驟卡：{num, title, body, photo, color}"""
    slots = [
        (Inches(0.42), left),
        (Inches(6.88), right),
    ]
    for x, info in slots:
        rect(slide, x, Inches(1.18), Inches(6.04), Inches(5.78), WHITE, rounded=True, adj=0.06)
        badge(slide, x + Inches(0.22), Inches(1.36), f"步驟 {info['num']}", info["color"])
        add_textbox(slide, x + Inches(1.50), Inches(1.34), Inches(4.3), Inches(0.46),
                    info["title"], 18, INK, True, FONT_TITLE, PP_ALIGN.LEFT, MSO_ANCHOR.MIDDLE)
        photo(slide, info["photo"], x + Inches(0.22), Inches(1.92),
              Inches(5.60), Inches(3.18), info["color"])
        add_textbox(slide, x + Inches(0.28), Inches(5.22), Inches(5.50), Inches(1.52),
                    info["body"], 15, INK, False, FONT_BODY, PP_ALIGN.LEFT, MSO_ANCHOR.TOP)


def chip_row(slide, items, y=Inches(1.28)):
    x = Inches(0.48)
    for text, fill, fg in items:
        w = Inches(0.28 + 0.32 * max(4, len(text) * 0.55))
        w = min(w, Inches(3.4))
        pill(slide, x, y, w, Inches(0.40), text, fill, fg)
        x += w + Inches(0.14)


def build():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    slides = []

    def S(bg=BG):
        s = new_slide(prs)
        add_bg(s, bg)
        slides.append(s)
        return s

    # ---------- 1 封面 ----------
    s = S()
    oval(s, Inches(5.4), Inches(-1.2), Inches(8.2), Inches(8.2), RGBColor(0xFF, 0xE3, 0xF0))
    add_textbox(s, Inches(0.7), Inches(0.55), Inches(7.2), Inches(0.4),
                "國中生活科技／家政課  ·  動手做甜點", 14, PINK, True)
    add_textbox(s, Inches(0.7), Inches(1.05), Inches(8.0), Inches(1.7),
                "今天做兩款\n超可愛中式甜點！", 40, INK, True, FONT_TITLE)
    add_textbox(s, Inches(0.7), Inches(2.85), Inches(7.6), Inches(0.7),
                "青花瓷山藥糕  +  雪梨菊花銀耳凍\n看圖就能做，步驟簡單、顏色超開心。", 18, INK_SOFT, False)

    pill(s, Inches(0.7), Inches(3.85), Inches(2.15), Inches(0.42), "免烤箱", SKY, WHITE)
    pill(s, Inches(3.00), Inches(3.85), Inches(2.35), Inches(0.42), "冷藏就成型", MINT, WHITE)
    pill(s, Inches(5.50), Inches(3.85), Inches(2.15), Inches(0.42), "好看又好吃", ORANGE, WHITE)

    photo(s, "qh_hero.jpg", Inches(8.35), Inches(0.85), Inches(4.45), Inches(2.55), BLUE)
    photo(s, "xl_hero.jpg", Inches(8.35), Inches(3.60), Inches(4.45), Inches(2.55), YELLOW)

    add_textbox(s, Inches(0.7), Inches(6.55), Inches(7.4), Inches(0.5),
                "畫面來自原影片，照著做就對了。", 12, INK_SOFT)

    # ---------- 2 兩款預告 ----------
    s = S()
    add_header(s, "先來看今天的兩位主角", PINK)
    rect(s, Inches(0.42), Inches(1.22), Inches(6.04), Inches(5.55), WHITE, rounded=True, adj=0.06)
    rect(s, Inches(6.88), Inches(1.22), Inches(6.04), Inches(5.55), WHITE, rounded=True, adj=0.06)
    badge(s, Inches(0.64), Inches(1.40), "第一款", BLUE, Inches(1.2))
    badge(s, Inches(7.10), Inches(1.40), "第二款", ORANGE, Inches(1.2))
    add_textbox(s, Inches(1.95), Inches(1.38), Inches(4.2), Inches(0.46),
                "青花瓷山藥糕", 20, INK, True, FONT_TITLE, PP_ALIGN.LEFT, MSO_ANCHOR.MIDDLE)
    add_textbox(s, Inches(8.42), Inches(1.38), Inches(4.2), Inches(0.46),
                "雪梨菊花銀耳凍", 20, INK, True, FONT_TITLE, PP_ALIGN.LEFT, MSO_ANCHOR.MIDDLE)
    photo(s, "qh_done.jpg", Inches(0.64), Inches(2.00), Inches(5.60), Inches(3.15), BLUE)
    photo(s, "xl_stack.jpg", Inches(7.10), Inches(2.00), Inches(5.60), Inches(3.15), YELLOW)
    add_textbox(s, Inches(0.64), Inches(5.28), Inches(5.60), Inches(1.2),
                "藍白雲紋包在透明凍裡，像一塊可以吃的小青花瓷。山藥香糯，外層 Q 彈。",
                15, INK)
    add_textbox(s, Inches(7.10), Inches(5.28), Inches(5.60), Inches(1.2),
                "下面是雪梨銀耳凍，上面再放一顆藏著菊花的水晶球。清甜、潤喉、超漂亮。",
                15, INK)

    # ---------- 3 安全 ----------
    s = S()
    add_header(s, "動手前，先記住這 4 件事", ORANGE, "安全小叮嚀")
    tips = [
        ("01", CORAL, "刀子請小心", "削皮、切塊時手指要遠離刀刃。不熟的同學請老師幫忙。"),
        ("02", ORANGE, "熱的東西會燙", "蒸箱、煮沸的湯和模具都燙手。戴隔熱手套，不要急著摸。"),
        ("03", SKY, "建議戴手套", "山藥黏手又可能讓皮膚癢，揉麵團時請戴手套。"),
        ("04", MINT, "白涼粉要煮沸", "不煮開就不會好好凝固。攪拌到粉完全化開，再倒模。"),
    ]
    for i, (num, color, title, body) in enumerate(tips):
        col, row = i % 2, i // 2
        x = Inches(0.42 + col * 6.46)
        y = Inches(1.22 + row * 2.70)
        rect(s, x, y, Inches(6.22), Inches(2.48), WHITE, rounded=True, adj=0.08)
        oval(s, x + Inches(0.28), y + Inches(0.38), Inches(0.78), Inches(0.78), color)
        add_textbox(s, x + Inches(0.28), y + Inches(0.38), Inches(0.78), Inches(0.78),
                    num, 16, WHITE, True, FONT_TITLE, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
        add_textbox(s, x + Inches(1.22), y + Inches(0.42), Inches(4.6), Inches(0.55),
                    title, 22, INK, True, FONT_TITLE)
        add_textbox(s, x + Inches(1.22), y + Inches(1.05), Inches(4.6), Inches(1.1),
                    body, 16, INK)

    # ---------- 4 青花瓷材料 ----------
    s = S(QH_BG)
    add_header(s, "青花瓷山藥糕｜先把材料排好", BLUE, "第一款")
    photo(s, "qh_yam.jpg", Inches(0.42), Inches(1.22), Inches(6.15), Inches(3.45), BLUE)
    photo(s, "qh_recipe.jpg", Inches(0.42), Inches(4.80), Inches(6.15), Inches(2.00), SKY)
    rect(s, Inches(6.80), Inches(1.22), Inches(6.10), Inches(5.58), WHITE, rounded=True, adj=0.06)
    add_textbox(s, Inches(7.05), Inches(1.38), Inches(5.6), Inches(0.45),
                "山藥團材料", 18, BLUE, True, FONT_TITLE)
    mats = [
        ("鐵棍山藥", "320 g", "選這個口感更粉糯"),
        ("熟糯米粉", "25 g", "讓它能搓成團"),
        ("白砂糖", "25 g", "甜度剛剛好"),
        ("蝶豆花粉", "適量", "染出青花藍"),
    ]
    y = Inches(1.90)
    for name, amt, note in mats:
        oval(s, Inches(7.12), y + Inches(0.08), Inches(0.22), Inches(0.22), BLUE)
        add_textbox(s, Inches(7.45), y, Inches(2.4), Inches(0.38), name, 15, INK, True)
        add_textbox(s, Inches(9.85), y, Inches(1.2), Inches(0.38), amt, 15, BLUE, True)
        add_textbox(s, Inches(11.05), y, Inches(1.6), Inches(0.38), note, 12, INK_SOFT)
        y += Inches(0.48)
    add_textbox(s, Inches(7.05), Inches(4.20), Inches(5.6), Inches(1.5),
                "透明外衣：椰子水 + 白涼粉（約 10：1）\n可包喜歡的餡，也可以不包，直接搓球。",
                14, INK)

    # ---------- 5 青花 1-2 ----------
    s = S(QH_BG)
    add_header(s, "青花瓷山藥糕｜去皮、蒸熟", BLUE, "第一款")
    step_pair(s, {
        "num": "1", "color": BLUE, "title": "選鐵棍山藥，削皮",
        "photo": "qh_peel.jpg",
        "body": "做糕點要選鐵棍山藥，口感更好。戴手套削皮，黏液比較不會弄到手。",
    }, {
        "num": "2", "color": SKY, "title": "放進蒸烤箱蒸熟",
        "photo": "qh_steam.jpg",
        "body": "去皮後切段或切小塊，放入蒸烤箱蒸到筷子能輕易插入。熱熱的更好壓泥。",
    })

    # ---------- 6 青花 3-4 ----------
    s = S(QH_BG)
    add_header(s, "青花瓷山藥糕｜壓泥、揉團", BLUE, "第一款")
    step_pair(s, {
        "num": "3", "color": LAVENDER, "title": "趁熱壓成細泥",
        "photo": "qh_mash.jpg",
        "body": "蒸熟取出，用壓泥器壓到細細滑滑、沒有顆粒。越細，成品越像瓷器。",
    }, {
        "num": "4", "color": PINK, "title": "加粉揉成團",
        "photo": "qh_knead.jpg",
        "body": "加入熟糯米粉 25g、白砂糖 25g，揉到不黏手的光滑麵團。",
    })

    # ---------- 7 青花 5-6 ----------
    s = S(QH_BG)
    add_header(s, "青花瓷山藥糕｜染青花藍、包餡搓球", BLUE, "第一款")
    step_pair(s, {
        "num": "5", "color": BLUE, "title": "一份加蝶豆花粉",
        "photo": "qh_blue.jpg",
        "body": "麵團分成兩份。其中一份加少量蝶豆花粉，揉成漂亮的藍色。粉少少加，顏色才清透。",
    }, {
        "num": "6", "color": CORAL, "title": "雙色混合，可包餡",
        "photo": "qh_fill.jpg",
        "body": "藍白兩色輕輕捏在一起，不要揉太勻，才有雲紋。可包喜歡的餡，再搓成圓球。",
    })

    # ---------- 8 青花 7-8 ----------
    s = S(QH_BG)
    add_header(s, "青花瓷山藥糕｜做透明外衣、入模", BLUE, "第一款")
    step_pair(s, {
        "num": "7", "color": MINT, "title": "椰子水加白涼粉",
        "photo": "qh_whisk.jpg",
        "body": "椰子水加白涼粉攪勻，煮沸。比例大約 10：1，例如 150 毫升椰子水配 15 克白涼粉。",
    }, {
        "num": "8", "color": SKY, "title": "倒模 → 放山藥團 → 再倒滿",
        "photo": "qh_set.jpg",
        "body": "先倒一層進模具，稍稍定型後放入山藥團，再倒滿。冷藏定型，脫模就完成了！",
    })

    # ---------- 9 青花完成 ----------
    s = S(QH_BG)
    add_header(s, "完成！可以吃的小青花瓷", BLUE, "第一款")
    photo(s, "qh_done.jpg", Inches(0.42), Inches(1.22), Inches(8.35), Inches(4.70), BLUE)
    rect(s, Inches(9.00), Inches(1.22), Inches(3.90), Inches(4.70), WHITE, rounded=True, adj=0.08)
    add_textbox(s, Inches(9.22), Inches(1.45), Inches(3.5), Inches(0.5),
                "成功密技", 18, BLUE, True, FONT_TITLE)
    secrets = [
        "藍白不要揉太開，紋路才像青花。",
        "外衣要等底部稍凝，球才不會沉到底。",
        "脫模前再冷藏一下，花紋更清楚。",
        "現做現吃最嫩，也可以冰著當點心。",
    ]
    box = s.shapes.add_textbox(Inches(9.22), Inches(2.05), Inches(3.5), Inches(3.5))
    tf = box.text_frame
    tf.word_wrap = True
    for i, line in enumerate(secrets):
        write_para(tf, "•  " + line, 14, INK, first=(i == 0), space_after=12)
    add_textbox(s, Inches(0.48), Inches(6.05), Inches(12.4), Inches(0.85),
                "當傳統遇上創新，這方寸之間的瓷意糕點，就是我們今天的第一份作品。",
                16, INK_SOFT)

    # ---------- 10 雪梨材料 ----------
    s = S(XL_BG)
    add_header(s, "雪梨菊花銀耳凍｜材料超簡單", ORANGE, "第二款")
    photo(s, "xl_hero.jpg", Inches(0.42), Inches(1.22), Inches(6.15), Inches(5.58), YELLOW)
    rect(s, Inches(6.80), Inches(1.22), Inches(6.10), Inches(5.58), WHITE, rounded=True, adj=0.06)
    add_textbox(s, Inches(7.05), Inches(1.42), Inches(5.6), Inches(0.4),
                "分成兩層來做", 18, ORANGE, True, FONT_TITLE)

    add_textbox(s, Inches(7.05), Inches(1.95), Inches(5.6), Inches(0.36),
                "雪梨銀耳凍", 16, INK, True)
    add_textbox(s, Inches(7.05), Inches(2.32), Inches(5.6), Inches(1.15),
                "雪梨銀耳湯 200 g\n白涼粉 20 g\n（先把雪梨、銀耳、冰糖煮成湯）", 15, INK)

    add_textbox(s, Inches(7.05), Inches(3.55), Inches(5.6), Inches(0.36),
                "菊花凍", 16, INK, True)
    add_textbox(s, Inches(7.05), Inches(3.92), Inches(5.6), Inches(1.0),
                "菊花茶 100 g\n白涼粉 10 g\n中間再放一朵泡好的菊花", 15, INK)

    rect(s, Inches(7.05), Inches(5.15), Inches(5.6), Inches(1.35), RGBColor(0xFF, 0xF3, 0xD1), rounded=True, adj=0.12)
    add_textbox(s, Inches(7.22), Inches(5.28), Inches(5.25), Inches(1.1),
                "好記口訣：湯／茶 : 白涼粉 ＝ 10 : 1\n這樣凍起來 Q 彈又不會太硬。",
                14, INK)

    # ---------- 11 雪梨 1-2 ----------
    s = S(XL_BG)
    add_header(s, "雪梨菊花銀耳凍｜泡茶、處理雪梨", ORANGE, "第二款")
    step_pair(s, {
        "num": "1", "color": YELLOW, "title": "先泡一壺菊花茶",
        "photo": "xl_bloom.jpg",
        "body": "準備菊花，熱水沖泡備用。待會一半拿來做凍，一朵漂亮的花要藏進圓球裡。",
    }, {
        "num": "2", "color": ORANGE, "title": "雪梨去皮切小塊",
        "photo": "xl_cut.jpg",
        "body": "雪梨削皮、去核，切成均勻小丁。塊不要太大，煮起來更快，凍裡也比較好看。",
    })

    # ---------- 12 雪梨 3-4 ----------
    s = S(XL_BG)
    add_header(s, "雪梨菊花銀耳凍｜煮一鍋潤潤的湯", ORANGE, "第二款")
    step_pair(s, {
        "num": "3", "color": MINT, "title": "下鍋，加銀耳悶煮",
        "photo": "xl_fungus.jpg",
        "body": "雪梨丁放進熱水裡，再加入銀耳。蓋上蓋子悶煮大約半小時，讓銀耳出膠、湯變黏潤。",
    }, {
        "num": "4", "color": CORAL, "title": "最後加冰糖煮沸",
        "photo": "xl_sugar.jpg",
        "body": "銀耳軟了以後，加入冰糖攪到溶化並煮沸。這鍋湯又香又甜，先聞一下超幸福。",
    })

    # ---------- 13 雪梨 5-6 ----------
    s = S(XL_BG)
    add_header(s, "雪梨菊花銀耳凍｜做成凍、倒進模具", ORANGE, "第二款")
    step_pair(s, {
        "num": "5", "color": SKY, "title": "取湯加白涼粉",
        "photo": "xl_powder.jpg",
        "body": "舀出雪梨銀耳湯 200 克，加入白涼粉 20 克，攪勻煮沸。粉要完全化開，才不會一顆顆。",
    }, {
        "num": "6", "color": BLUE, "title": "倒入模具冷藏",
        "photo": "xl_mold.jpg",
        "body": "煮沸後趁熱舀進花形模具，連同雪梨塊和銀耳一起放。送進冰箱冷藏定型。",
    })

    # ---------- 14 雪梨 7-8 ----------
    s = S(XL_BG)
    add_header(s, "雪梨菊花銀耳凍｜做會開花的水晶球", ORANGE, "第二款")
    step_pair(s, {
        "num": "7", "color": LAVENDER, "title": "菊花茶加白涼粉煮沸",
        "photo": "xl_jupowder.jpg",
        "body": "另起一鍋，倒入泡好的菊花茶 100 克，加白涼粉 10 克，攪勻煮沸。顏色會淡淡金黃。",
    }, {
        "num": "8", "color": PINK, "title": "倒模，中間放一朵花",
        "photo": "xl_flowerin.jpg",
        "body": "倒進圓球模具，用夾子放進一朵泡好的菊花，再補滿茶凍。冷藏定型就完成了！",
    })

    # ---------- 15 雪梨完成 ----------
    s = S(XL_BG)
    add_header(s, "疊在一起，像一朵會融化的秋日", ORANGE, "第二款")
    photo(s, "xl_base.jpg", Inches(0.42), Inches(1.22), Inches(6.04), Inches(3.40), MINT)
    photo(s, "xl_stack.jpg", Inches(6.68), Inches(1.22), Inches(6.22), Inches(3.40), YELLOW)
    rect(s, Inches(0.42), Inches(4.80), Inches(12.48), Inches(1.95), WHITE, rounded=True, adj=0.08)
    add_textbox(s, Inches(0.70), Inches(4.95), Inches(12.0), Inches(0.4),
                "怎麼擺盤？", 16, ORANGE, True, FONT_TITLE)
    add_textbox(s, Inches(0.70), Inches(5.38), Inches(12.0), Inches(1.15),
                "脫模後，先放上雪梨銀耳凍當底座，再把菊花水晶球輕輕放上去。想更漂亮，可以撒一點菊花碎、或連湯一起盛在小碗裡。銀耳的膠質鎖住雪梨的清甜，再搭配菊花凍的清雅——這不僅是味覺的享受。",
                15, INK)

    # ---------- 16 對照複習 ----------
    s = S()
    add_header(s, "兩款對照，上課超好記", PINK, "快速複習")
    # left card
    rect(s, Inches(0.42), Inches(1.20), Inches(6.04), Inches(5.55), WHITE, rounded=True, adj=0.06)
    rect(s, Inches(0.42), Inches(1.20), Inches(6.04), Inches(0.62), BLUE, rounded=True, adj=0.18)
    add_textbox(s, Inches(0.42), Inches(1.20), Inches(6.04), Inches(0.62),
                "青花瓷山藥糕", 18, WHITE, True, FONT_TITLE, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    qh_lines = [
        "1  鐵棍山藥去皮蒸熟",
        "2  壓泥，加粉揉團",
        "3  一份染成蝶豆花藍",
        "4  雙色雲紋搓球（可包餡）",
        "5  椰子凍倒模當外衣",
        "6  冷藏脫模，完成青花瓷",
    ]
    box = s.shapes.add_textbox(Inches(0.75), Inches(2.05), Inches(5.4), Inches(4.4))
    tf = box.text_frame
    tf.word_wrap = True
    for i, line in enumerate(qh_lines):
        write_para(tf, line, 18, INK, first=(i == 0), space_after=10)

    rect(s, Inches(6.88), Inches(1.20), Inches(6.04), Inches(5.55), WHITE, rounded=True, adj=0.06)
    rect(s, Inches(6.88), Inches(1.20), Inches(6.04), Inches(0.62), ORANGE, rounded=True, adj=0.18)
    add_textbox(s, Inches(6.88), Inches(1.20), Inches(6.04), Inches(0.62),
                "雪梨菊花銀耳凍", 18, WHITE, True, FONT_TITLE, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    xl_lines = [
        "1  先泡菊花茶備用",
        "2  雪梨切丁，加銀耳悶煮",
        "3  冰糖煮沸成湯",
        "4  湯＋白涼粉 → 雪梨凍",
        "5  茶＋白涼粉＋菊花 → 花球",
        "6  冷藏後疊在一起擺盤",
    ]
    box = s.shapes.add_textbox(Inches(7.22), Inches(2.05), Inches(5.4), Inches(4.4))
    tf = box.text_frame
    tf.word_wrap = True
    for i, line in enumerate(xl_lines):
        write_para(tf, line, 18, INK, first=(i == 0), space_after=10)

    # ---------- 17 結尾 ----------
    s = S()
    oval(s, Inches(3.8), Inches(0.2), Inches(5.8), Inches(5.8), RGBColor(0xFF, 0xE8, 0xF2))
    add_textbox(s, Inches(0.8), Inches(1.35), Inches(11.7), Inches(1.3),
                "今天你學會了兩件事：", 28, INK, True, FONT_TITLE, PP_ALIGN.CENTER)
    add_textbox(s, Inches(1.4), Inches(2.7), Inches(10.5), Inches(1.6),
                "把山藥變成可以吃的青花瓷，\n把一壺秋日甜湯變成會開花的果凍。",
                22, INK, False, FONT_TITLE, PP_ALIGN.CENTER)
    pill(s, Inches(3.15), Inches(4.55), Inches(3.15), Inches(0.55), "動手就會成功", PINK, WHITE)
    pill(s, Inches(6.55), Inches(4.55), Inches(3.55), Inches(0.55), "做好記得拍照打卡", YELLOW, INK)
    add_textbox(s, Inches(1.2), Inches(5.45), Inches(10.9), Inches(1.1),
                "白涼粉記得煮沸、模具記得冷藏。\n下一堂課，我們再做別的快樂甜點！",
                16, INK_SOFT, False, FONT_BODY, PP_ALIGN.CENTER)

    total = len(slides)
    notes = [""] * total
    notes[0] = "中式甜點動手做  ·  給國中同學的快樂食譜"
    for i, sl in enumerate(slides, start=1):
        add_footer(sl, i, total, notes[0])

    prs.save(OUT)
    print(f"saved {OUT}  ({total} slides)")


if __name__ == "__main__":
    build()
